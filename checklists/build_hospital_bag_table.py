#!/usr/bin/env python3
"""Generate shareable hospital-bag tables (xlsx + html) from one data source."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT_DIR = Path(__file__).resolve().parent

HEADERS = ["分类", "物品", "建议数量", "谁准备", "何时带/何时买", "是否已备", "备注"]

HOSPITAL_ROWS = [
    ["证件袋", "产妇身份证（原件）", "1", "产妇/家属", "入院当天随身", "", "复印件另夹一份"],
    ["证件袋", "配偶身份证（原件）", "1", "家属", "入院当天随身", "", "办出生证明、落户也会用"],
    ["证件袋", "医保电子凭证 + 社保卡", "已登录即可", "产妇", "入院当天随身", "", "「湘医保」提前登录，别只靠现场网络"],
    ["证件袋", "医院诊疗卡 / 就诊卡", "1", "产妇", "入院当天随身", "", "有卡就带，没有用身份证办理"],
    ["证件袋", "产检手册和全部报告", "全套", "产妇", "入院当天随身", "", "含 B 超、化验、糖耐、胎监"],
    ["证件袋", "住院证", "1", "产妇", "医生开了立刻放进袋", "", "没住院证不能办住院"],
    ["证件袋", "结婚证、户口本", "原件", "家属", "办证时再送，不必整天放病房", "", "出生证明、落户常用"],
    ["证件袋", "银行卡 + 少量现金", "2 张卡 / 少量现金", "家属", "入院当天", "", "预交款常见 2000–3000 元，以窗口为准"],
    ["证件袋", "血型 / 过敏 / 紧急联系人纸条", "1 张", "家属", "入院当天随身", "", "写预产期、血型、过敏史、谁接电话"],
    ["产妇·衣服", "待产裙或哺乳衣（前开扣）", "2–3 套", "产妇", "36 周前备好", "", "住院几天换洗用"],
    ["产妇·衣服", "纯棉内裤（一次性或普通）", "8–10 条", "产妇", "36 周前备好", "", "用量比平时月经期多"],
    ["产妇·衣服", "防滑拖鞋 + 袜子", "各 1–2", "产妇", "36 周前备好", "", "医院地面常湿"],
    ["产妇·衣服", "薄外套", "1 件", "产妇", "36 周前备好", "", "产房/病房空调可能凉；秋冬再加一件"],
    ["产妇·衣服", "出院换洗衣物", "1 套", "家属", "出院当天送来即可", "", "不必提前堆在病房"],
    ["产妇·护理", "产妇卫生巾（日用+夜用）", "够 3–5 天", "产妇", "36 周前备好", "", "量要比平时月经多"],
    ["产妇·护理", "产褥垫 / 看护垫", "8–10 张", "产妇", "36 周前备好", "", "垫床用"],
    ["产妇·护理", "防溢乳垫 + 乳头舒缓膏", "少量", "产妇", "36 周前备好", "", "住院开奶用，不必屯一大箱"],
    ["产妇·护理", "湿巾、纸巾、垃圾袋", "各 1 包", "产妇", "36 周前备好", "", "病房垃圾自己收比较方便"],
    ["产妇·护理", "唇膏、免洗洗手液", "各 1", "产妇", "入院小件", "", "病房偏干，洗手不便时用"],
    ["产妇·护理", "会阴冲洗器或挤瓶", "1", "家属", "顺产后可再补", "", "按护士习惯，入院后再买也行"],
    ["产妇·护理", "腹带", "1", "家属", "剖宫产再备", "", "是否用、松紧听医生护士，不要自己勒太紧"],
    ["产妇·洗漱用品", "牙刷牙膏、毛巾、洗面奶、梳子", "个人用量", "产妇", "36 周前备好", "", "护肤品少量即可"],
    ["产妇·洗漱用品", "带盖杯 + 吸管", "1 套", "产妇", "36 周前备好", "", "剖宫产早期喝水方便"],
    ["产妇·洗漱用品", "充电器、充电宝、耳机", "各 1", "产妇/家属", "入院当天", "", "充电宝满电"],
    ["产妇·洗漱用品", "框架眼镜", "1", "产妇", "如平时戴隐形则改框架", "", "住院少戴隐形更省事"],
    ["产妇·吃的", "少量零食和自己习惯的水", "1 小袋", "家属", "入院当天", "", "苏打饼干、独立包装坚果即可；医院有食堂"],
    ["宝宝（3–5天）", "连体衣 / 和尚服（纯棉无骨缝）", "3–5 件", "家属", "36 周前备好", "", "医院通常有部分被服，仍建议自备"],
    ["宝宝（3–5天）", "包被 / 包巾", "2 条", "家属", "36 周前备好", "", "秋冬加厚一床，病房仍可能开空调"],
    ["宝宝（3–5天）", "帽子、袜子、护手", "各 1–2", "家属", "36 周前备好", "", ""],
    ["宝宝（3–5天）", "纸尿裤（新生儿码）", "1 小包", "家属", "32–34 周买", "", "先买小包装，出院按实际用量补"],
    ["宝宝（3–5天）", "隔尿垫", "若干", "家属", "36 周前备好", "", ""],
    ["宝宝（3–5天）", "湿巾（无酒精无香精）、柔纸巾", "各 1 包", "家属", "36 周前备好", "", "棉签按护士指导用，不要自行挖鼻耳"],
    ["家属陪护", "家属身份证", "1", "家属", "入院当天", "", "办陪护证"],
    ["家属陪护", "换洗衣物、拖鞋、洗漱", "1–2 天用量", "家属", "入院当天", "", "夜间薄毯或外套一件"],
    ["家属陪护", "充电宝、水杯", "各 1", "家属", "入院当天", "", ""],
    ["家属陪护", "食堂/外卖支付方式", "已开通", "家属", "入院前确认", "", "微信/支付宝/银行卡能付即可"],
    ["家属陪护", "陪护证", "入院后办理", "家属", "办好随身带", "", "进产房着装以护士要求为准"],
]

DONT_BRING_ROWS = [
    ["不要带", "热得快、电煮锅、电炉", "—", "全员", "全程不要", "", "住院须知：病房不用明火和电热器具"],
    ["不要带", "大量奶粉、奶瓶", "—", "全员", "先不带", "", "先按医院喂养指导；需要配方奶再按医嘱买"],
    ["不要带", "贵重首饰、大量现金", "—", "全员", "全程不要", "", "贵重物品自己保管，病房地方小也不安全"],
    ["不要带", "过多换季衣物、一年用量纸尿裤", "—", "全员", "不要一次搬进医院", "", "只带几天用量，家里再备一个月的"],
]

HOME_ROWS = [
    ["家里·产妇", "产褥期卫生巾、内裤", "够回家后一周再补", "家属", "出院前家里备好", "", "不必全部搬进医院"],
    ["家里·产妇", "哺乳内衣", "3–4 件", "产妇", "36 周前后", "", ""],
    ["家里·产妇", "吸奶器", "按需", "产妇", "住院后再决定买/租", "", "不是人人必须，避免闲置"],
    ["家里·产妇", "收腹带、盆底训练用品", "按需", "产妇", "42 天复查后再听康复意见", "", "不要自行勒太紧"],
    ["家里·婴儿", "婴儿床或安全同房睡方案", "1", "家属", "出院前", "", "不建议同被捂睡"],
    ["家里·婴儿", "床单、睡袋或包被", "够换洗", "家属", "出院前", "", "室温大约 24–26℃ 作参考"],
    ["家里·婴儿", "夜灯、尿布台或床边整理筐、垃圾桶", "各 1", "家属", "出院前", "", ""],
    ["家里·婴儿", "洗澡盆、浴巾、水温计", "1 套", "家属", "出院前", "", ""],
    ["家里·婴儿", "婴儿指甲剪、电子体温计", "各 1", "家属", "出院前", "", ""],
    ["家里·安全", "窗户限位、电线收纳、热水瓶放高处", "按家情况", "家属", "出院前", "", "满月前减少无关探视，洗手再抱"],
    ["家里·小药箱", "体温计、棉签、无菌纱布", "家用少量", "家属", "出院前", "", "不替代就医；不要自行给新生儿用药"],
]

META = {
    "title": "待产包清单（可发给家人一起勾选）",
    "hospital": "目标医院：长沙市中心医院（南华大学附属长沙中心医院）",
    "address": "地址：长沙市雨花区韶山南路 161 号（铁道学院对面）",
    "deadline": "建议完成：孕 36 周前封箱，放门口或后备箱（夜间也能 5 分钟拿走）",
    "note": "原则：医院只带证件 + 几天内用得上的；家里再备回来一个月要用的。医疗决策以产科当面医嘱为准。",
}

THIN = Border(
    left=Side(style="thin", color="D4C4B0"),
    right=Side(style="thin", color="D4C4B0"),
    top=Side(style="thin", color="D4C4B0"),
    bottom=Side(style="thin", color="D4C4B0"),
)

CAT_COLORS = {
    "证件袋": "FDE8D0",
    "产妇·衣服": "F8D7DA",
    "产妇·护理": "F8D7DA",
    "产妇·洗漱用品": "F8D7DA",
    "产妇·吃的": "F8D7DA",
    "宝宝（3–5天）": "D4EDDA",
    "家属陪护": "D6EAF8",
    "不要带": "F5C6CB",
    "家里·产妇": "FFF3CD",
    "家里·婴儿": "D4EDDA",
    "家里·安全": "E2E3E5",
    "家里·小药箱": "E2E3E5",
}


def style_header_row(ws, row: int, cols: int) -> None:
    fill = PatternFill("solid", fgColor="7A4E2D")
    font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def write_data_rows(ws, start_row: int, rows: list[list[str]]) -> None:
    body = Font(name="Microsoft YaHei", size=10)
    for i, row in enumerate(rows):
        r = start_row + i
        cat_fill = PatternFill("solid", fgColor=CAT_COLORS.get(row[0], "FFFFFF"))
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.font = body
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c == 1:
                cell.fill = cat_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif c == 6:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.value = "☐"
                cell.font = Font(name="Microsoft YaHei", size=14)
            else:
                cell.fill = PatternFill("solid", fgColor="FFFCF7")


def decorate_sheet(ws, title_extra: str, rows: list[list[str]], table_name: str) -> None:
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    ws.print_title_rows = "1:6"
    ws.freeze_panes = "A7"
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[6].height = 22
    ws.merge_cells("A1:G1")
    ws.merge_cells("A2:G2")
    ws.merge_cells("A3:G3")
    ws.merge_cells("A4:G4")
    ws.merge_cells("A5:G5")
    ws["A1"] = META["title"] + title_extra
    ws["A1"].font = Font(name="Microsoft YaHei", bold=True, size=16, color="5C3317")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for row, key in enumerate(["hospital", "address", "deadline", "note"], start=2):
        ws.cell(row, 1, META[key])
        ws.cell(row, 1).font = Font(name="Microsoft YaHei", size=10, color="5C3317")
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 20 if row < 5 else 32
    for col, header in enumerate(HEADERS, 1):
        ws.cell(6, col, header)
    style_header_row(ws, 6, 7)
    write_data_rows(ws, 7, rows)
    last = 6 + len(rows)
    widths = [16, 32, 18, 12, 22, 12, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(7, last + 1):
        ws.row_dimensions[r].height = 32
    table = Table(displayName=table_name, ref=f"A6:G{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    ws.add_table(table)
    ws.auto_filter.ref = f"A6:G{last}"
    ws.oddFooter.left.text = "家庭备产清单 · 不能替代医嘱"
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"


def write_xlsx(path: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "入院待产包"
    decorate_sheet(ws1, "", HOSPITAL_ROWS, "HospitalBag")
    ws2 = wb.create_sheet("不要带")
    decorate_sheet(ws2, " — 不要带进病房", DONT_BRING_ROWS, "DoNotBring")
    ws3 = wb.create_sheet("家里准备")
    decorate_sheet(ws3, " — 不必搬进医院", HOME_ROWS, "HomePrep")
    wb.save(path)


def html_table(rows: list[list[str]]) -> str:
    parts = [
        "<table>",
        "<thead><tr>"
        + "".join(f"<th>{h}</th>" for h in HEADERS)
        + "</tr></thead><tbody>",
    ]
    for row in rows:
        cat = row[0]
        slug = {
            "证件袋": "id",
            "产妇·衣服": "mom",
            "产妇·护理": "mom",
            "产妇·洗漱用品": "mom",
            "产妇·吃的": "mom",
            "宝宝（3–5天）": "baby",
            "家属陪护": "family",
            "不要带": "no",
            "家里·产妇": "home",
            "家里·婴儿": "baby",
            "家里·安全": "home",
            "家里·小药箱": "home",
        }.get(cat, "id")
        cells = "".join(
            f'<td class="check">☐</td>' if i == 5 else f"<td>{c}</td>"
            for i, c in enumerate(row)
        )
        parts.append(f'<tr class="{slug}">{cells}</tr>')
    parts.append("</tbody></table>")
    return "\n".join(parts)


def write_html(path: Path) -> None:
    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>待产包清单</title>
  <style>
    :root {{ --ink:#5c3317; --line:#e6d5c3; }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0; padding: 24px;
      font-family: "PingFang SC", "Microsoft YaHei", sans-serif;
      color: #2b2118; background: #faf6f1;
    }}
    .sheet {{
      max-width: 1100px; margin: 0 auto; background: #fff;
      border: 1px solid var(--line); padding: 28px 24px 36px;
    }}
    h1 {{ margin: 0 0 8px; text-align: center; color: var(--ink); font-size: 24px; }}
    .meta {{ text-align: center; color: #6b5344; font-size: 14px; line-height: 1.7; margin: 0; }}
    h2 {{ color: var(--ink); font-size: 18px; margin: 28px 0 10px; border-left: 4px solid #c4783a; padding-left: 10px; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
    th, td {{ border: 1px solid var(--line); padding: 8px 10px; vertical-align: top; }}
    th {{ background: #7a4e2d; color: #fff; font-weight: 600; white-space: nowrap; }}
    td.check {{ text-align: center; font-size: 18px; width: 64px; }}
    tr.id td:first-child {{ background: #fde8d0; }}
    tr.mom td:first-child {{ background: #f8d7da; }}
    tr.baby td:first-child {{ background: #d4edda; }}
    tr.family td:first-child {{ background: #d6eaf8; }}
    tr.no td:first-child {{ background: #f5c6cb; }}
    tr.home td:first-child {{ background: #fff3cd; }}
    .hint {{ margin-top: 20px; font-size: 13px; color: #6b5344; }}
    @media print {{
      body {{ background: #fff; padding: 0; }}
      .sheet {{ border: none; padding: 0; max-width: none; }}
      h2 {{ break-after: avoid; }}
      table {{ break-inside: auto; }}
      tr {{ break-inside: avoid; }}
    }}
  </style>
</head>
<body>
  <div class="sheet">
    <h1>{META["title"]}</h1>
    <p class="meta">{META["hospital"]}<br>{META["address"]}<br>{META["deadline"]}<br>{META["note"]}</p>
    <h2>一、入院待产包</h2>
    {html_table(HOSPITAL_ROWS)}
    <h2>二、不要带进病房</h2>
    {html_table(DONT_BRING_ROWS)}
    <h2>三、家里准备（不必搬进医院）</h2>
    {html_table(HOME_ROWS)}
    <p class="hint">发给别人时：直接发本 HTML，或用浏览器打开后「打印 → 另存为 PDF」。Excel 版见同目录「待产包清单.xlsx」。勾选列打印后可用笔勾。</p>
  </div>
</body>
</html>
"""
    path.write_text(html, encoding="utf-8")


def write_csv(path: Path, rows: list[list[str]]) -> None:
    import csv

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["说明", META["title"]])
        w.writerow(["医院", META["hospital"]])
        w.writerow(["地址", META["address"]])
        w.writerow(["完成时间", META["deadline"]])
        w.writerow([])
        w.writerow(HEADERS)
        for row in rows:
            out = list(row)
            out[5] = ""
            w.writerow(out)


def main() -> None:
    write_xlsx(OUT_DIR / "待产包清单.xlsx")
    write_html(OUT_DIR / "待产包清单.html")
    write_csv(OUT_DIR / "待产包清单-入院.csv", HOSPITAL_ROWS + DONT_BRING_ROWS)
    write_csv(OUT_DIR / "待产包清单-家里.csv", HOME_ROWS)
    print("wrote xlsx/html/csv")


if __name__ == "__main__":
    main()
